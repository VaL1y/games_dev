from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("ТЗ_Лабиринт_с_костями_v0.1.docx")

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF4D6"
GOLD = "8A6500"
PALE_RED = "FDECEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "202124"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, size=11, color=BLACK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style(style, size, color, before, after, bold=True, line=1.0):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = True


def add_numbering_definition(doc, num_id, abstract_id, fmt, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    apply_numbering(p, 41)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_step(doc, text):
    p = doc.add_paragraph()
    apply_numbering(p, 42)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=DARK_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), fill)
        borders.append(border)
    p_pr.append(borders)
    r = p.add_run(label + " ")
    set_run_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, 120)
    set_table_borders(table)
    for idx, label in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=9.2, color=BLACK)
    set_table_geometry(table, widths, 120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

set_style(doc.styles["Heading 1"], 16, BLUE, 18, 10)
set_style(doc.styles["Heading 2"], 13, BLUE, 14, 7)
set_style(doc.styles["Heading 3"], 12, DARK_BLUE, 10, 5)

add_numbering_definition(doc, 41, 41, "bullet", "•")
add_numbering_definition(doc, 42, 42, "decimal", "%1.")

# Running furniture: compact-reference-guide + memo masthead, without a bottom border.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("ЛАБИРИНТ С КОСТЯМИ  /  ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
set_run_font(hr, size=8.5, color=MUTED, bold=True)
footer = section.footer
add_page_number(footer.paragraphs[0])

# Cover / memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
set_run_font(r, size=12, color=GOLD, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Лабиринт с костями")
set_run_font(r, size=28, color=NAVY, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Пошаговая настольная гонка с ловушками, ставкой на победу и преследующими тенями")
set_run_font(r, size=13.5, color=MUTED)

meta = [
    ("Версия", "0.1 — концепт и трехэтапный план"),
    ("Назначение", "Основа для прототипирования, оценки и последующей разработки"),
    ("Приоритет", "Этап 1: минимальный играбельный MVP"),
    ("Статус правил", "Рабочая гипотеза; спорные места вынесены в раздел открытых решений"),
]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + ": ")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(value)
    set_run_font(r, size=10.5, color=BLACK)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
add_callout(
    doc,
    "Коротко.",
    "Игроки бросают две кости и идут к финишу первого мира. Первый дошедший решает: зафиксировать победу или рискнуть наградой и перейти во второй, более опасный мир. Позже за рискнувшим начинает охоту ИИ-тень; у игрока есть фора в 6 ходов.",
)

doc.add_heading("1. Цель документа", level=1)
add_body(doc, "Документ превращает исходную идею в проверяемый объем разработки. Этап 1 обязателен и должен дать короткую, но законченную игровую сессию. Этапы 2 и 3 опциональны: к ним стоит переходить только после проверки интереса к основной развилке «остановиться или рискнуть». ")

doc.add_heading("2. Концепция продукта", level=1)
add_body(doc, "Жанр: пошаговая соревновательная настольная игра для 1–4 участников, где движение определяется двумя костями, клетки дают очки и события, а финиш первого поля не обязательно завершает партию.")
add_bullet(doc, "Главная эмоция: напряжение перед выбором — сохранить уже полученную победу или поставить ее на кон ради большей награды.", "Главная эмоция:")
add_bullet(doc, "Главная отличительная механика: второй мир и тени, которые превращают отстающих или завершивших путь игроков в угрозу для лидера.", "Главная отличительная механика:")
add_bullet(doc, "Базовая длительность сессии: MVP — 8–15 минут; расширенная версия — 20–35 минут.", "Базовая длительность сессии:")
add_bullet(doc, "Целевая платформа и технология пока не зафиксированы. ТЗ описывает игровую логику независимо от движка.", "Целевая платформа и технология:")

doc.add_heading("3. Термины", level=1)
add_body(doc, "Первый мир — основное поле-гонка от старта до точки Б. Второй мир — отдельное более сложное поле для игроков, принявших риск. Тень — управляемый ИИ преследователь. Безопасный финиш — отказ от риска с фиксацией результата. Абсолютная победа — успешное прохождение второго мира.")

doc.add_heading("4. Игровой цикл", level=1)
for step in [
    "Игрок видит поле, список участников в правом нижнем углу и блок костей в левом верхнем углу.",
    "В свой ход игрок бросает две шестигранные кости и перемещается на сумму значений.",
    "После остановки начисляются очки клетки и разыгрывается ее эффект: обычная клетка, ловушка или случайное событие.",
    "После достижения финиша первого мира игрок выбирает безопасную победу либо переход во второй мир.",
    "Рискнувшие проходят усиленное поле; при выполнении условия появления за ними стартует тень с форой игроку в 6 его ходов.",
    "Партия заканчивается абсолютной победой во втором мире либо запасным итогом первого мира, если никто не смог пройти риск-маршрут.",
]:
    add_step(doc, step)

doc.add_heading("5. Сводка трех этапов", level=1)
add_table(
    doc,
    ["Этап", "Задача", "Что должно быть видно игроку", "Решение после этапа"],
    [
        ("1. MVP", "Проверить основной цикл и выбор риска", "Одна партия, 2 поля, кости, 3 типа клеток, простая тень", "Интересно ли повторять партию?"),
        ("2. Расширение", "Сделать полноценную реиграбельную игру", "Разные события, комнаты, боты, прогресс, баланс", "Работает ли удержание и соревнование?"),
        ("3. Полная концепция", "Раскрыть асимметрию теней и метаигру жизней", "Глубокий второй мир, месть, редкие монеты, контент", "Есть ли смысл развивать как сервис?"),
    ],
    [1250, 2250, 3360, 2500],
)

doc.add_page_break()
doc.add_heading("ЭТАП 1. Минимальный MVP", level=1)
add_callout(doc, "Цель этапа.", "За одну короткую сессию показать движение по клеткам, ловушки, начисление очков, развилку на финише и погоню тени. Визуальная простота допустима; непонятные правила и незавершаемая партия — нет.")

doc.add_heading("1.1. Объем MVP", level=2)
add_bullet(doc, "Режим: локальная партия на одном устройстве; 1 человек и до 3 простых ботов либо 2–4 человека по очереди.")
add_bullet(doc, "Первый мир: одно линейное поле на 24 клетки, старт и финиш; без развилок маршрута.")
add_bullet(doc, "Второй мир: отдельная дорожка на 12 клеток с более частыми негативными событиями.")
add_bullet(doc, "Кости: 2d6; ход равен сумме. Для финиша не требуется точный бросок.")
add_bullet(doc, "Типы клеток: обычная, ловушка, случайное событие, финиш.")
add_bullet(doc, "Партия полностью проходит от стартового экрана до экрана результатов без ручного перезапуска состояния.")

doc.add_heading("1.2. Экран партии", level=2)
add_bullet(doc, "Центр: игровое поле, фишки, номер каждой клетки и выделение текущего игрока.")
add_bullet(doc, "Левый верхний угол: две кости, кнопка «Бросить», сумма и короткая подпись результата.")
add_bullet(doc, "Правый нижний угол: список игроков — цвет, имя, место, очки, состояние и число пропущенных ходов.")
add_bullet(doc, "Нижняя/центральная область: журнал последних 5 событий, чтобы результат броска и ловушки можно было проверить.")
add_bullet(doc, "Модальное окно на финише: «Забрать победу» и «Рискнуть ради большей награды» с понятным описанием последствий.")

doc.add_heading("1.3. Правила клеток MVP", level=2)
add_table(
    doc,
    ["Тип", "Частота", "Эффект"],
    [
        ("Обычная", "16 из 24", "Только начисление очков за ход."),
        ("Ловушка", "4 из 24", "Либо пропуск следующего хода, либо возврат на 3 клетки."),
        ("Случайная", "3 из 24", "Один из эффектов: +200 очков, -200 очков, вперед/назад на 2 клетки."),
        ("Финиш", "1 из 24", "Начисление финишного бонуса и выбор безопасного или рискованного пути."),
    ],
    [1800, 1700, 5860],
)

doc.add_heading("1.4. Начисление очков", level=2)
add_body(doc, "Рабочая формула исходной идеи: очки за ход = номер клетки, на которой игрок закончил движение, × сумма двух костей. Пример: остановка на клетке 15 при броске 6 и 2 дает 15 × 8 = 120 очков.")
add_table(
    doc,
    ["Событие", "MVP-значение", "Комментарий"],
    [
        ("Первым достичь финиша первого мира", "+2 000", "Начисляется один раз за партию."),
        ("Пройти второй мир и выжить", "+5 000", "Награда абсолютному победителю."),
        ("Остановить соперника своей тенью", "+500", "В MVP — за захват соперника тенью."),
        ("Никого не остановить", "0 бонуса", "Сгорает только бонус за остановку, не все набранные очки."),
    ],
    [3600, 1700, 4060],
)
add_callout(doc, "Баланс.", "Формула быстро увеличивает награду ближе к финишу. Для MVP это допустимо как тест исходной идеи; после первых 20–30 партий нужно проверить, не обесценивает ли она ранние ходы.", fill=PALE_GOLD, label_color=GOLD)

doc.add_heading("1.5. Выбор на финише", level=2)
add_body(doc, "Безопасный выбор: игрок фиксирует результат первого мира, больше не делает ходов и наблюдает за партией. Если никто не завершит второй мир, лучший безопасный результат определяется по порядку достижения финиша, затем по очкам.")
add_body(doc, "Рискованный выбор: игрок сохраняет текущие очки, но откладывает окончательную победу и переносится во второй мир. При провале риск-маршрута бонус +5 000 не начисляется; уже заработанные очки не сгорают в MVP.")

doc.add_heading("1.6. Тень и фора в 6 ходов", level=2)
add_body(doc, "Рабочая интерпретация для прототипа: первый рискнувший начинает второй мир без преследователя. Когда следующий участник достигает финиша первого мира, в начале второго мира появляется тень, которая преследует первого рискнувшего. Счетчик форы равен 6 ходам преследуемого игрока; до его обнуления тень не двигается.")
add_bullet(doc, "После форы тень бросает 1d6 после каждого хода цели и движется по тому же маршруту.")
add_bullet(doc, "Если тень оказывается на клетке цели или проходит дальше, цель считается пойманной и завершает риск-маршрут неудачей.")
add_bullet(doc, "В MVP тень не выбирает маршрут и не использует способности; это визуально понятный таймер давления.")

doc.add_heading("1.7. Условия завершения MVP", level=2)
add_bullet(doc, "Абсолютная победа: первый участник дошел до конца второго мира раньше тени.")
add_bullet(doc, "Запасной победитель: если второй мир никто не прошел, выигрывает первый игрок с зафиксированным безопасным финишем; если таких нет — участник с наибольшими очками.")
add_bullet(doc, "Экран результатов показывает победителя, тип победы, итоговые очки, бонусы и ключевые события партии.")

doc.add_heading("1.8. Поведение ботов", level=2)
add_bullet(doc, "Бот автоматически бросает кости и применяет эффект клетки с паузой 0,5–1,0 секунды для читаемости.")
add_bullet(doc, "На первом тестовом профиле бот выбирает риск с вероятностью 50%; вероятность задается настройкой.")
add_bullet(doc, "Бот не должен зависать, совершать ход дважды или принимать решение за человека.")

doc.add_heading("1.9. Что не входит в MVP", level=2)
add_bullet(doc, "Онлайн-мультиплеер, аккаунты, чат, подбор соперников и переподключение.")
add_bullet(doc, "Магазин, покупка жизни, редкие монеты как расходуемая валюта и механика мести.")
add_bullet(doc, "Редактор карт, косметика, сюжет, сезонные события, рейтинги и достижения.")
add_bullet(doc, "Сложный ИИ теней, разные классы персонажей и развилки маршрутов.")

doc.add_heading("1.10. Критерии приемки MVP", level=2)
for item in [
    "Новую партию можно начать за три действия или меньше.",
    "Полная партия с ботами гарантированно завершается; отсутствуют бесконечные состояния.",
    "Каждый бросок, перемещение, начисление очков и эффект клетки отражаются в журнале.",
    "Оба выбора на финише приводят к разным и понятным состояниям игры.",
    "Фора тени уменьшается только по ходам преследуемого игрока и равна ровно 6 ходам.",
    "Победитель и расчет очков воспроизводимы по журналу партии.",
    "Интерфейс корректно показывает 1–4 участников при типовом разрешении выбранной платформы.",
]:
    add_bullet(doc, "□ " + item)

doc.add_page_break()
doc.add_heading("ЭТАП 2. Расширенная основная игра", level=1)
add_callout(doc, "Условие старта.", "Переходить к этапу после подтверждения, что игроки понимают выбор риска и хотя бы часть тестеров хочет сыграть повторно.")

doc.add_heading("2.1. Цель", level=2)
add_body(doc, "Превратить демонстрационный прототип в реиграбельную соревновательную игру: добавить разнообразие первого мира, устойчивую сессию, понятный прогресс и базовые социальные режимы.")

doc.add_heading("2.2. Состав этапа", level=2)
add_bullet(doc, "2–4 игрока; одиночная игра с ботами, локальная партия и приватная онлайн-комната.")
add_bullet(doc, "Три карты первого мира по 30–45 клеток; на каждой — минимум одна развилка с коротким опасным и длинным безопасным путем.")
add_bullet(doc, "Не менее 8 ловушек и 12 случайных событий с защитой от длинных серий одинаковых эффектов.")
add_bullet(doc, "Статусы: щит от ловушки, ускорение, замедление и один переброс костей.")
add_bullet(doc, "Обучение первой партии, пауза, повтор партии, сохранение настроек, реконнект в приватную комнату.")
add_bullet(doc, "Подробный экран результатов: источник каждого бонуса, место, пройденный путь и история ключевых решений.")
add_bullet(doc, "Телеметрия баланса без персональных данных: длительность, частота риска, завершение второго мира, отрывы по очкам.")

doc.add_heading("2.3. Развитие правил", level=2)
add_bullet(doc, "Уточнить и зафиксировать владение тенью: чья тень появляется, кому начисляются +500 и что считается «остановил игрока».")
add_bullet(doc, "Ввести банк риск-награды: бонусы второго мира видны до выбора, но выдаются только при успешном выходе.")
add_bullet(doc, "Добавить ограничитель разгона очков: тестировать формулу клетки × бросок, нормализацию по кругу и потолок за ход.")
add_bullet(doc, "Исключить безнадежные партии: игрок, отставший из-за серии ловушек, получает редкое событие возвращения, но не гарантированную победу.")

doc.add_heading("2.4. Базовая метаигра", level=2)
add_bullet(doc, "Профиль игрока, статистика партий и простые достижения.")
add_bullet(doc, "Редкие монеты выдаются за игровые достижения и абсолютную победу; на этапе 2 их можно копить, но нельзя покупать за реальные деньги.")
add_bullet(doc, "Косметические награды допустимы, если они не меняют шанс победы.")

doc.add_heading("2.5. Критерии готовности этапа", level=2)
add_bullet(doc, "Не менее 80% тестовых онлайн-партий завершаются без разрыва состояния или ручного вмешательства.")
add_bullet(doc, "Игрок после обучения может объяснить разницу между безопасным и рискованным финишем.")
add_bullet(doc, "Нет одной стратегии выбора маршрута или риска, которая доминирует во всех тестовых конфигурациях.")
add_bullet(doc, "Все начисления очков и монет имеют проверяемый источник в итоговом протоколе.")

doc.add_heading("2.6. Необязательные дополнения", level=2)
add_bullet(doc, "Публичный подбор соперников и рейтинг.")
add_bullet(doc, "Эмоции/быстрые сообщения без свободного текстового чата.")
add_bullet(doc, "Набор косметических фишек и визуальных тем поля.")

doc.add_page_break()
doc.add_heading("ЭТАП 3. Полная концепция: второй мир, тени и жизни", level=1)
add_callout(doc, "Статус этапа.", "Опциональный продуктовый слой. Его объем следует подтверждать данными этапа 2, потому что он заметно увеличивает требования к балансу, ИИ и объяснению правил.")

doc.add_heading("3.1. Цель", level=2)
add_body(doc, "Сделать второй мир самостоятельной кульминацией партии, где лидер получает высокую награду, но действия остальных игроков материализуются в виде теней и угроз. Добавить долгосрочную валюту жизней и возможность ответного хода против соперника без гарантированной мести.")

doc.add_heading("3.2. Полноценный второй мир", level=2)
add_bullet(doc, "Отдельная визуальная палитра и читаемое ощущение более опасного пространства.")
add_bullet(doc, "Несколько маршрутов, клетки распада, рискованные сокращения и точки временного укрытия от тени.")
add_bullet(doc, "Каждый новый финишер первого мира создает событие во втором: новую тень, усиление существующей либо препятствие — по заранее видимому правилу.")
add_bullet(doc, "Фора первого рискнувшего остается 6 ходов, но модификаторы карты могут временно менять скорость тени, а не сам размер базовой форы.")
add_bullet(doc, "ИИ тени оценивает расстояние, сокращения и уязвимые состояния, однако его решения объяснимы через подсветку намерения.")

doc.add_heading("3.3. Взаимодействие игроков через тени", level=2)
add_bullet(doc, "Игрок, завершивший первый мир, не превращается в пассивного наблюдателя: он может один раз выбрать поведение связанной с ним тени из 2–3 понятных вариантов.")
add_bullet(doc, "Управление ограничено, чтобы исключить сговор и затягивание партии; тень остается преимущественно ИИ-сущностью.")
add_bullet(doc, "За остановку соперника начисляется +500 владельцу связанной тени. Если остановок нет, сгорает только этот потенциальный бонус.")

doc.add_heading("3.4. Редкие монеты и покупка жизни", level=2)
add_body(doc, "Редкие монеты — долгосрочная игровая валюта. Базовая безопасная версия механики: перед партией игрок может потратить ограниченное число монет на одну дополнительную жизнь. Жизнь позволяет один раз вернуться после захвата тенью или смертельной ловушки на последнюю безопасную клетку.")
add_bullet(doc, "Одна дополнительная жизнь максимум на партию; цену и частоту получения определяет баланс.")
add_bullet(doc, "Решение о покупке принимается до начала партии либо в явно обозначенной точке, а не после просмотра случайного результата.")
add_bullet(doc, "Для соревновательного рейтинга дополнительные жизни отключаются или выдаются всем на одинаковых условиях.")
add_bullet(doc, "Механика «отомстить противнику» реализуется как одноразовая контр-ловушка или усиление своей тени, а не как прямое исключение выбранного игрока.")

doc.add_heading("3.5. Награда абсолютному победителю", level=2)
add_bullet(doc, "Итоговые очки партии, +5 000 за выживание и заранее показанный риск-множитель/банк.")
add_bullet(doc, "Редкая монета либо прогресс к ней по прозрачному правилу.")
add_bullet(doc, "Косметический знак абсолютной победы в профиле и итоговом экране.")
add_bullet(doc, "Награда не должна давать бесконечное преимущество в следующих матчах.")

doc.add_heading("3.6. Критерии готовности этапа", level=2)
add_bullet(doc, "Второй мир создает отдельную тактическую задачу, а не только более длинную дорожку.")
add_bullet(doc, "Игрок до подтверждения риска понимает, что может получить и что потерять.")
add_bullet(doc, "Поведение тени прогнозируемо по интерфейсу, но не сводится к гарантированному исходу.")
add_bullet(doc, "Покупка жизни не дает платного или накопительного доминирования.")
add_bullet(doc, "Партия имеет максимальную длительность и механизм принудительного завершения при затягивании.")

doc.add_page_break()
doc.add_heading("6. Состояния игрока и партии", level=1)
add_table(
    doc,
    ["Состояние", "Что означает", "Разрешенные действия"],
    [
        ("Ожидает ход", "Игрок активен в текущем мире", "Просмотр поля и журнала"),
        ("Ходит", "Текущий активный игрок", "Бросок, выбор эффекта, завершение хода"),
        ("Пропускает ход", "Действует ловушка/статус", "Автоматический пропуск с пояснением"),
        ("Безопасный финиш", "Результат первого мира зафиксирован", "Наблюдение; на поздних этапах — ограниченная связь с тенью"),
        ("Во втором мире", "Игрок принял риск", "Ход, применение предметов, уклонение от тени"),
        ("Пойман", "Риск-маршрут провален", "Наблюдение либо трата заранее купленной жизни"),
        ("Завершил", "Получен окончательный результат", "Просмотр итогов и повтор партии"),
    ],
    [1900, 3460, 4000],
)

doc.add_heading("7. Нефункциональные требования", level=1)
add_bullet(doc, "Детерминированность: при одинаковом seed и одинаковых решениях партия воспроизводится для отладки.")
add_bullet(doc, "Сохранность состояния: переход между мирами, выбор риска и появление тени выполняются атомарно.")
add_bullet(doc, "Производительность: плавная анимация на целевой платформе; анимации не блокируют логику и могут быть ускорены.")
add_bullet(doc, "Доступность: цвет игрока дублируется номером/символом; важные состояния не кодируются только цветом.")
add_bullet(doc, "Локализация: весь пользовательский текст вынесен из кода; первая обязательная локаль — русский язык.")
add_bullet(doc, "Журналирование: ошибки хода, рассинхронизация, начисления и смена состояния партии доступны разработчику.")

doc.add_heading("8. Минимальная модель данных", level=1)
add_bullet(doc, "GameSession: идентификатор, seed, этап/мир, номер раунда, очередь, статус и победитель.")
add_bullet(doc, "PlayerState: имя, цвет/символ, тип управления, позиция, очки, статусы, выбранный путь, жизни.")
add_bullet(doc, "BoardCell: индекс, тип, параметры эффекта и визуальный идентификатор.")
add_bullet(doc, "TurnRecord: игрок, значения костей, исходная/конечная позиция, эффект, изменение очков.")
add_bullet(doc, "ShadowState: владелец, цель, позиция, остаток форы, скорость, статус.")
add_bullet(doc, "RewardLedger: базовые очки, финишные бонусы, остановки, редкие монеты и причины начисления.")

doc.add_heading("9. Открытые продуктовые решения", level=1)
add_callout(doc, "Важно.", "Ниже перечислены вопросы, которых недостаточно в исходном описании. В MVP приняты безопасные рабочие ответы; перед этапами 2–3 их нужно подтвердить владельцу идеи.", fill=PALE_RED, label_color=RED)
questions = [
    ("Кто именно создает тень?", "MVP: следующий финишер запускает тень против первого рискнувшего."),
    ("Чья это тень?", "MVP: нейтральный ИИ; на этапе 3 она связывается с завершившим игроком."),
    ("Что сгорает, если никто никого не остановил?", "MVP: только потенциальные +500, не общий счет."),
    ("Что теряет рискнувший при провале?", "MVP: +5 000 и абсолютную победу; базовые очки сохраняются."),
    ("Могут ли несколько игроков идти во второй мир?", "MVP: да; абсолютным победителем становится первый завершивший."),
    ("Когда и как покупается жизнь?", "Не входит в MVP; базовая версия этапа 3 — до партии, одна жизнь максимум."),
    ("Что означает «отомстить»?", "Этап 3: ограниченная контр-ловушка/усиление тени, без прямого удаления соперника."),
    ("Какая целевая платформа?", "Решить до реализации интерфейса: браузер, ПК или мобильное устройство."),
]
add_table(doc, ["Вопрос", "Рабочее решение"], questions, [3650, 5710], header_fill=LIGHT_GRAY)

doc.add_heading("10. Порядок разработки и проверки", level=1)
for step in [
    "Зафиксировать платформу и управление; собрать серый интерактивный макет поля.",
    "Реализовать детерминированное ядро правил без анимаций и покрыть тестами движение, очки, финиш, тень и окончание партии.",
    "Подключить интерфейс MVP, журнал событий и простых ботов.",
    "Провести 20–30 наблюдаемых партий; измерить длительность, понимание выбора, частоту риска и причины досрочного выхода.",
    "Исправить правила, затем принять решение о переходе к этапу 2. Этап 3 не начинать до подтверждения ценности второго мира.",
]:
    add_step(doc, step)

doc.add_heading("11. Итоговый состав поставки по этапам", level=1)
add_body(doc, "Этап 1 — один законченный прототип с одной картой каждого мира, локальной партией, ботами, подсчетом очков и результатами. Этап 2 — реиграбельная основная игра с несколькими картами, онлайн-комнатой, прогрессом и балансировкой. Этап 3 — глубокий второй мир, управляемые правила появления теней, редкие монеты, жизни и ограниченная механика мести.")
add_body(doc, "Главный критерий успеха проекта: игроку интересно не просто первым дойти до точки Б, а осознанно решить, достаточно ли ему обычной победы или стоит рискнуть всем ради абсолютной.", italic=True)

doc.core_properties.title = "Техническое задание — Лабиринт с костями"
doc.core_properties.subject = "Трехэтапный план разработки игры"
doc.core_properties.author = "Проектная команда"
doc.core_properties.keywords = "игра, MVP, кости, лабиринт, тени, техническое задание"

doc.save(OUT)
print(OUT)
